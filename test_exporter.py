import unittest
from io import BytesIO
from zipfile import ZipFile
import pymupdf
from openpyxl import Workbook
from exporter import addresses, export_zip, render_pdf

class ExportTests(unittest.TestCase):
    def test_ranges(self):
        self.assertEqual(addresses('A1:B2，A1 C3'), ['A1', 'B1', 'A2', 'B2', 'C3'])
        for spec in ['A0', 'XFE1', 'A1:A5001', 'B2:A1', '']:
            with self.assertRaises(ValueError):
                addresses(spec)

    def test_pdf_content_size_and_fit(self):
        data = render_pdf('中文测试\nHello 123', 100, 60)
        with pymupdf.open(stream=data, filetype='pdf') as doc:
            self.assertEqual(len(doc), 1)
            self.assertAlmostEqual(doc[0].rect.width, 100*72/25.4, places=3)
            self.assertAlmostEqual(doc[0].rect.height, 60*72/25.4, places=3)
            self.assertIn('中文测试', doc[0].get_text())
            self.assertIn('Hello 123', doc[0].get_text())
            self.assertTrue(doc.extract_font(doc[0].get_fonts()[0][0])[3])
        render_pdf('自动换行测试'*20, 50, 30)
        with self.assertRaises(ValueError):
            render_pdf('太长'*1000, 10, 10, margin_mm=1)
        with self.assertRaises(ValueError):
            render_pdf('test', 10, 10, margin_mm=5)

    def test_zip_cells_and_blanks(self):
        book = Workbook()
        book.active.title = '标签'
        book.active['A1'] = '你好'
        book.active['A2'] = 0
        book.create_sheet('Second')['A1'] = False
        data, count, skipped = export_zip(book, book.sheetnames, 'A1:A3', width_mm=80, height_mm=40)
        self.assertEqual((count, skipped), (3, 3))
        with ZipFile(BytesIO(data)) as archive:
            self.assertEqual(archive.namelist(), ['01_标签/A1.pdf', '01_标签/A2.pdf', '02_Second/A1.pdf'])
            for filename in archive.namelist():
                with pymupdf.open(stream=archive.read(filename), filetype='pdf') as doc:
                    self.assertEqual(len(doc), 1)
        _, count, skipped = export_zip(book, ['标签'], 'A3', skip_empty=False, width_mm=80, height_mm=40)
        self.assertEqual((count, skipped), (1, 0))

if __name__ == '__main__':
    unittest.main()
