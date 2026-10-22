from io import BytesIO
import hashlib
import streamlit as st
from openpyxl import load_workbook
from exporter import addresses, export_zip

st.set_page_config(page_title='单元格转 PDF', page_icon='📄', layout='centered')
st.title('Excel 单元格 → PDF')
st.caption('批量选择 · 自定义尺寸 · 每格一页一文件')
st.write('上传 Excel，选定单元格或区域，即可下载包含独立 PDF 的 ZIP 压缩包。')
upload = st.file_uploader('上传 Excel 文件（支持 .xlsx / .xlsm）', type=['xlsx', 'xlsm'])
st.info('导出单元格的文字、数字和换行，不复制 Excel 的字体、颜色、图片或边框。公式使用 Excel 上次保存的计算结果；请先在 Excel 中重新计算并保存。旧版 .xls 请另存为 .xlsx。')
if upload is None:
    st.stop()
try:
    raw = upload.getvalue()
    book = load_workbook(BytesIO(raw), data_only=True)
except Exception:
    st.error('无法读取文件，请确认它是有效且未加密的 Excel 工作簿。')
    st.stop()

sheets = st.multiselect('选择工作表', book.sheetnames, default=book.sheetnames[:1])
spec = st.text_input('需要导出的单元格 / 区域', 'A1:A10', help='例如 A1, C3, D2:D8；多个区域用逗号、空格或换行分隔。每张选中工作表使用同一选择。合并单元格仅左上角含值。')
with st.expander('预览选择的内容', expanded=True):
    try:
        cells = addresses(spec)
        if sheets:
            st.dataframe([{'工作表': name, '单元格': cell,
                           '内容': str(book[name][cell].value) if book[name][cell].value is not None else ''}
                          for name in sheets[:3] for cell in cells[:20]], hide_index=True)
            st.caption('预览最多显示前 3 张表各 20 个单元格。')
    except ValueError as exc:
        st.warning(str(exc))

st.subheader('PDF 页面设置')
a, b = st.columns(2)
width = a.number_input('宽度（毫米）', min_value=10.0, max_value=1000.0, value=100.0)
height = b.number_input('高度（毫米）', min_value=10.0, max_value=1000.0, value=60.0)
margin = a.number_input('四周边距（毫米）', min_value=0.0, max_value=499.0, value=5.0)
font_size = b.number_input('字号（pt）', min_value=4, max_value=200, value=18)
align = a.selectbox('水平对齐', ['居中', '左对齐', '右对齐'])
vertical = b.selectbox('垂直对齐', ['居中', '顶部', '底部'])
skip = st.checkbox('跳过空单元格（含没有缓存结果的公式）', value=True)
st.caption('文字自动换行，内容过长时自动缩小字号，最低 4pt；仍无法容纳时会指出对应单元格。每个 PDF 固定只有一页。')
options = dict(width_mm=width, height_mm=height, margin_mm=margin, font_size=font_size,
               align={'居中': 'center', '左对齐': 'left', '右对齐': 'right'}[align],
               vertical={'居中': 'middle', '顶部': 'top', '底部': 'bottom'}[vertical])
key = hashlib.sha256(raw + repr((sheets, spec, skip, options)).encode()).hexdigest()
if st.button('生成 PDF 压缩包', type='primary', disabled=not sheets):
    st.session_state.pop('export', None)
    try:
        with st.spinner('正在生成 PDF…'):
            data, count, skipped = export_zip(book, sheets, spec, skip_empty=skip, **options)
        if count:
            st.session_state['export'] = (key, data, count, skipped)
        else:
            st.warning('没有可导出的内容，请检查选择范围，或取消“跳过空单元格”。')
    except Exception as exc:
        st.error(f'生成失败：{exc}')
result = st.session_state.get('export')
if result and result[0] == key:
    _, data, count, skipped = result
    st.success(f'已生成 {count} 个 PDF，跳过 {skipped} 个空单元格。')
    st.download_button('下载全部 PDF（ZIP）', data, file_name='cell_pdfs.zip', mime='application/zip', type='primary')
